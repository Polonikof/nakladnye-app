# -*- coding: utf-8 -*-
"""Проверка packing list Юсуф №1643: конвертация, распознавание, обучение."""
from __future__ import print_function

import os
import shutil
import tempfile
import unittest

import convert_nakladnaya as core
import learn_novye

HERE = os.path.dirname(os.path.abspath(__file__))
SAMPLES = os.path.join(HERE, "samples")
SRC = os.path.join(SAMPLES, "Юсуф исходник №1643.xlsx")
GOT = os.path.join(SAMPLES, "Готовый Юсуф №1643.xlsx")


class YusufConvertTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not os.path.isfile(SRC) or not os.path.isfile(GOT):
            raise unittest.SkipTest("нет samples/Юсуф …1643.xlsx")

    def test_detect_source_and_ready(self):
        self.assertEqual(core.opredelit_tip_faila(SRC), "yusuf")
        self.assertTrue(core.eto_yusuf_syroj(SRC))
        self.assertFalse(core.eto_yusuf_gotovyj(SRC))
        self.assertTrue(core.eto_yusuf_gotovyj(GOT))
        self.assertFalse(core.eto_yusuf_syroj(GOT))
        self.assertEqual(core.opredelit_tip_faila(GOT), "skip")
        self.assertFalse(core.eto_oboi_gotovyj(GOT))
        self.assertFalse(core.eto_oboi_syroj(SRC))

    def test_detect_packing_list_without_yusuf_in_name(self):
        tmp = tempfile.mkdtemp(prefix="yusuf_name_")
        try:
            anon = os.path.join(tmp, "invoice_1643.xlsx")
            shutil.copy2(SRC, anon)
            self.assertTrue(core.eto_packing_list(anon))
            self.assertEqual(core.opredelit_tip_faila(anon), "yusuf")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_convert_matches_ready(self):
        tmp = tempfile.mkdtemp(prefix="yusuf_out_")
        try:
            out = os.path.join(tmp, "out.xlsx")
            rows, path = core.convert_yusuf(SRC, out)
            self.assertEqual(rows, 89)
            self.assertEqual(path, out)
            diffs = core.sravnit_yusuf(out, GOT)
            self.assertEqual(diffs, [], msg="\n".join(diffs[:20]))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_name_building(self):
        extra = core.yusuf_dop_polya({
            "collection": "ETNO",
            "design": "9624A",
            "color": "BURGUNDY / BURGUNDY",
            "width": 150,
            "length": 230,
            "size_shape": "150 x 230 D",
        })
        self.assertEqual(extra["naim"], "ETNO 1,5*2,3 Прямоугольник")
        self.assertEqual(extra["full"], "Ковер ETNO 1,5*2,3 Прямоугольник")
        self.assertEqual(extra["char"], "9624A BURGUNDY/BURGUNDY")
        self.assertEqual(extra["color"], "BURGUNDY/BURGUNDY")
        self.assertEqual(extra["width_cm"], 150)
        self.assertEqual(extra["letter"], "D")
        extra2 = core.yusuf_dop_polya({
            "collection": "ETNO",
            "design": "9624A",
            "color": "BURGUNDY / BURGUNDY",
            "width": 80,
            "length": 150,
            "size_shape": "80 x 150 D",
        })
        self.assertEqual(extra2["naim"], "ETNO 0,8*1,5 Прямоугольник")

    def test_process_folder_sees_yusuf(self):
        tmp = tempfile.mkdtemp(prefix="yusuf_folder_")
        try:
            shutil.copy2(SRC, os.path.join(tmp, "Юсуф исходник №1643.xlsx"))
            sources, already, unknown = core.find_source_files(tmp, state={})
            self.assertEqual(unknown, [])
            self.assertEqual(already, [])
            self.assertEqual(len(sources), 1)
            self.assertEqual(sources[0][1], "yusuf")
            itog = core.process_folder(tmp, force=True)
            self.assertEqual(itog["errors"], [])
            self.assertEqual(len(itog["done"]), 1)
            self.assertEqual(itog["done"][0]["type"], "yusuf")
            self.assertEqual(itog["rows"], 89)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_learn_pair(self):
        inbox = tempfile.mkdtemp(prefix="yusuf_learn_")
        catalog = tempfile.mkdtemp(prefix="yusuf_cat_")
        try:
            shutil.copy2(SRC, os.path.join(inbox, "Юсуф исходник №1643.xlsx"))
            shutil.copy2(GOT, os.path.join(inbox, "Готовый Юсуф №1643.xlsx"))
            pairs = learn_novye.naiti_pary(inbox)
            self.assertEqual(len(pairs), 1, pairs)
            self.assertEqual(pairs[0]["kind"], "yusuf")
            itog = learn_novye.process_inbox(
                inbox, force=True, catalog_folder=catalog)
            self.assertEqual(itog["errors"], [])
            kinds = [x.get("type") for x in itog["done"]]
            self.assertIn("learn-yusuf", kinds)
            self.assertTrue(os.path.isfile(
                os.path.join(catalog, core.YUSUF_PROFILE_FILENAME)))
            reports = itog.get("reports") or []
            self.assertTrue(reports)
            self.assertEqual(reports[0].get("diffs"), [])
        finally:
            shutil.rmtree(inbox, ignore_errors=True)
            shutil.rmtree(catalog, ignore_errors=True)

    def test_skip_if_gotovyj_unchanged(self):
        inbox = tempfile.mkdtemp(prefix="yusuf_skip_")
        catalog = tempfile.mkdtemp(prefix="yusuf_cat_")
        try:
            shutil.copy2(SRC, os.path.join(inbox, "Юсуф исходник №1643.xlsx"))
            shutil.copy2(GOT, os.path.join(inbox, "Готовый Юсуф №1643.xlsx"))
            first = learn_novye.process_inbox(
                inbox, force=False, catalog_folder=catalog)
            self.assertTrue(first["reports"])
            second = learn_novye.process_inbox(
                inbox, force=False, catalog_folder=catalog)
            self.assertEqual(second["reports"], [])
            self.assertTrue(second["skipped"])
        finally:
            shutil.rmtree(inbox, ignore_errors=True)
            shutil.rmtree(catalog, ignore_errors=True)

    def test_relearn_overwrites_rules_when_gotovyj_changes(self):
        from openpyxl import load_workbook

        inbox = tempfile.mkdtemp(prefix="yusuf_relearn_")
        catalog = tempfile.mkdtemp(prefix="yusuf_cat_")
        try:
            src_p = os.path.join(inbox, "Юсуф исходник №1643.xlsx")
            got_p = os.path.join(inbox, "Готовый Юсуф №1643.xlsx")
            shutil.copy2(SRC, src_p)
            shutil.copy2(GOT, got_p)
            first = learn_novye.process_inbox(
                inbox, force=False, catalog_folder=catalog)
            self.assertTrue(first["reports"])
            prof = core.load_yusuf_profile(catalog)
            self.assertEqual(prof.get("unit"), "шт")
            self.assertEqual(prof.get("shape_letters", {}).get("D"), "Прямоугольник")

            wb = load_workbook(got_p, data_only=True)
            ws = wb.active
            patched = []
            for r in range(20, 109):
                patched.append((
                    r,
                    str(ws.cell(r, 15).value or ""),
                    str(ws.cell(r, 17).value or ""),
                ))
            wb.close()
            wb = load_workbook(got_p)
            ws = wb.active
            for r, naim, full in patched:
                ws.cell(r, 15).value = naim.replace("Прямоугольник", "прямоуг.")
                ws.cell(r, 17).value = full.replace("Прямоугольник", "прямоуг.")
                ws.cell(r, 16).value = "шт."
            wb.save(got_p)
            wb.close()

            second = learn_novye.process_inbox(
                inbox, force=False, catalog_folder=catalog)
            self.assertTrue(second["reports"], "должны переобучить, готовый изменился")
            self.assertEqual(second.get("skipped"), [])
            prof = core.load_yusuf_profile(catalog)
            self.assertEqual(prof.get("unit"), "шт.")
            self.assertEqual(prof.get("shape_letters", {}).get("D"), "прямоуг.")

            extra = core.yusuf_dop_polya({
                "collection": "ETNO",
                "design": "9624A",
                "color": "BURGUNDY / BURGUNDY",
                "width": 200,
                "length": 300,
                "size_shape": "200 x 300 D",
            }, profile=prof)
            self.assertEqual(extra["unit"], "шт.")
            self.assertEqual(extra["naim"], "ETNO 2*3 прямоуг.")
            self.assertEqual(extra["full"], "Ковер ETNO 2*3 прямоуг.")
        finally:
            shutil.rmtree(inbox, ignore_errors=True)
            shutil.rmtree(catalog, ignore_errors=True)


class EmployeePackTests(unittest.TestCase):
    def test_zip_contains_exe_and_catalogs(self):
        catalog = tempfile.mkdtemp(prefix="pack_cat_")
        try:
            shutil.copy2(os.path.join(HERE, "oboi_catalog.json"),
                         os.path.join(catalog, core.OBOI_CATALOG_FILENAME))
            shutil.copy2(os.path.join(HERE, "yusuf_profile.json"),
                         os.path.join(catalog, core.YUSUF_PROFILE_FILENAME))
            fake_exe = os.path.join(catalog, "fake.exe")
            with open(fake_exe, "wb") as f:
                f.write(b"MZ-fake-nakladnye")
            info = learn_novye.sobrat_paket_dlya_sotrudnika(
                catalog, "1.5.1", fake_exe)
            self.assertTrue(os.path.isfile(info["zip"]))
            self.assertTrue(info["zip"].endswith("Накладные_1.5.1_для_сотрудника.zip"))
            import zipfile
            with zipfile.ZipFile(info["zip"], "r") as zf:
                names = set(zf.namelist())
            self.assertIn("Накладные.exe", names)
            self.assertIn("oboi_catalog.json", names)
            self.assertIn("yusuf_profile.json", names)
            self.assertIn("version.txt", names)
            self.assertIn("ПРОЧТИТЕ.txt", names)
            with zipfile.ZipFile(info["zip"], "r") as zf:
                self.assertEqual(zf.read("Накладные.exe"), b"MZ-fake-nakladnye")
        finally:
            shutil.rmtree(catalog, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
