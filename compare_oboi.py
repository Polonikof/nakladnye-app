# Сравнение результата скрипта (Обои) с эталоном «готовый»
#
#   python compare_oboi.py samples/out_oboi.xlsx "samples/ВВП УПД 04.09.26 готовый.xlsx"
#
import sys
from openpyxl import load_workbook


def art_key(v):
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip() if v is not None else ""


def main():
    if len(sys.argv) < 3:
        print("Usage: python compare_oboi.py <результат.xlsx> <эталон_готовый.xlsx>")
        sys.exit(2)

    out_path, got_path = sys.argv[1], sys.argv[2]
    out = load_workbook(out_path, data_only=True)["Обработанные данные"]
    got = load_workbook(got_path, data_only=True).worksheets[0]

    got_rows = []
    for r in range(1, got.max_row + 1):
        code = got.cell(r, 1).value
        if not (isinstance(code, str) and code.strip().upper().startswith("УТ")):
            continue
        got_rows.append({
            "code": str(code).strip(),
            "name": str(got.cell(r, 18).value).strip(),
            "art": art_key(got.cell(r, 19).value),
            "qty": str(got.cell(r, 36).value).strip(),
            "price": str(got.cell(r, 42).value).strip(),
            "sum_wo": str(got.cell(r, 49).value).strip(),
            "nds": str(got.cell(r, 68).value).strip(),
            "sum_w": str(got.cell(r, 76).value).strip(),
        })

    out_rows = []
    for r in range(2, out.max_row + 1):
        out_rows.append({
            "code": str(out.cell(r, 1).value or "").strip(),
            "name": str(out.cell(r, 3).value).strip(),
            "art": art_key(out.cell(r, 4).value),
            "qty": str(out.cell(r, 8).value).strip(),
            "price": str(out.cell(r, 9).value).strip(),
            "sum_wo": str(out.cell(r, 10).value).strip(),
            "nds": str(out.cell(r, 13).value).strip(),
            "sum_w": str(out.cell(r, 14).value).strip(),
        })

    print(f"результат: {len(out_rows)} стр., эталон: {len(got_rows)} стр.")
    if len(out_rows) != len(got_rows):
        print("РАСХОЖДЕНИЕ: разное число строк")
        sys.exit(1)

    fields = ["code", "name", "art", "qty", "price", "sum_wo", "nds", "sum_w"]
    diffs = 0
    for i, (o, g) in enumerate(zip(out_rows, got_rows), 1):
        bad = [f for f in fields if o[f] != g[f]]
        if bad:
            diffs += 1
            print(f"#{i}: " + "; ".join(f"{f}: {o[f]!r} != {g[f]!r}" for f in bad))

    if diffs:
        print(f"Расхождений: {diffs}")
        sys.exit(1)
    print("OK: 0 расхождений")


if __name__ == "__main__":
    main()
